"""Excel 紀錄檔的測試，包含防護 4（公式注入）。全部使用合成資料。"""

from __future__ import annotations

import os
import re
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ticket_agent import excel_store
from ticket_agent.excel_store import RECORD_HEADERS, RECORD_SHEET, STATS_HEADERS, STATS_SHEET, StoreError, TicketStore, looks_like_formula, write_cell
from ticket_agent.models import STATUS_ACTIVE, STATUS_DELETED, STATUS_REMOVED, TicketData

NOW = datetime(2026, 10, 24, 15, 30, 12)
UID_A = "123456789012345678"
UID_B = "987654321098765432"

PAYLOADS = [
    "=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1)", "\t=1+1", "\r=1+1", " =1+1", "\n=1+1", "＝1+1",
    '=HYPERLINK("http://example.com","x")', "=cmd|' /C calc'!A0", "@import", "  +1", "=IMPORTXML(\"http://x\",\"//a\")",
]


def make_store(tmp_path, clock=lambda: NOW) -> TicketStore:
    folder = tmp_path / "data"
    folder.mkdir(exist_ok=True)
    return TicketStore(folder / "tickets.xlsx", clock=clock)


def ticket(**overrides) -> TicketData:
    base = dict(ticket_type="演唱會", name="測試演唱會 2099", event_datetime="2099-12-31 19:30", location="假的場館",
                seat_or_type="A 區 12 排", amount=1200.0, currency="TWD", other_info="票號：TEST-0001")
    base.update(overrides)
    return TicketData(**base)


def add(store, *, uploader="小明", uid=UID_A, data=None, photo="photos/a.jpg", when=NOW, reply=None):
    return store.add_record(uploader=uploader, uploader_id=uid, data=data or ticket(), photo_path=photo,
                            recorded_at=when, reply_message_id=reply)


def sheet_xml(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/worksheets/")]


def has_formula(path: Path) -> bool:
    return any(re.search(r"<f[ >/]", xml) for xml in sheet_xml(path))


# ---------- 建檔與標題列 ----------

def test_creates_file_with_two_sheets_and_headers(tmp_path):
    store = make_store(tmp_path)
    wb = load_workbook(store.path)
    assert wb.sheetnames == [RECORD_SHEET, STATS_SHEET]
    assert tuple(c.value for c in wb[RECORD_SHEET][1]) == RECORD_HEADERS
    assert tuple(c.value for c in wb[STATS_SHEET][1]) == STATS_HEADERS
    assert (store.path.stat().st_mode & 0o777) == 0o600


def test_wrong_headers_refuse_to_start(tmp_path):
    store = make_store(tmp_path)
    wb = load_workbook(store.path)
    wb[RECORD_SHEET]["A1"] = "別人的欄位"
    wb.save(store.path)
    with pytest.raises(StoreError, match="標題列"):
        TicketStore(store.path)


def test_corrupted_file_raises_store_error_without_leaking_content(tmp_path):
    path = tmp_path / "data" / "tickets.xlsx"
    path.parent.mkdir()
    path.write_bytes(b"this is not an xlsx file, SECRET-CONTENT")
    with pytest.raises(StoreError) as info:
        TicketStore(path)
    assert "SECRET-CONTENT" not in str(info.value)


# ---------- 一直往下加，不覆蓋 ----------

def test_records_are_appended_not_overwritten_and_survive_restart(tmp_path):
    store = make_store(tmp_path)
    first = add(store)
    second = add(store, data=ticket(name="第二張"), photo="photos/b.jpg")
    reopened = TicketStore(store.path, clock=lambda: NOW)
    third = add(reopened, data=ticket(name="第三張"), photo="photos/c.jpg")
    assert [first.serial, second.serial, third.serial] == [1, 2, 3]
    names = [r.data.name for r in reopened.all_records()]
    assert names == ["測試演唱會 2099", "第二張", "第三張"]
    assert load_workbook(store.path)[RECORD_SHEET].max_row == 4  # 標題列 + 3 筆


def test_fields_round_trip_exactly(tmp_path):
    store = make_store(tmp_path)
    saved = add(store, data=ticket(amount=1234.5), photo="photos/a.jpg", reply="111222333444555666")
    store.merge_into(saved.serial, ticket(), "photos/newer.jpg")
    loaded = store.get(saved.serial)
    assert loaded.uploader_id == UID_A  # 18 位數字必須維持文字，不被轉成浮點數
    assert loaded.recorded_at == NOW
    assert loaded.data.amount == 1200.0  # 被新照片補齊成 1200
    assert loaded.photo_paths == ("photos/newer.jpg", "photos/a.jpg")
    assert loaded.reply_message_ids == ("111222333444555666",)
    assert loaded.status == STATUS_ACTIVE


def test_amount_is_stored_as_a_number_and_empty_as_blank(tmp_path):
    store = make_store(tmp_path)
    add(store)
    add(store, data=ticket(amount=None, currency=""))
    ws = load_workbook(store.path)[RECORD_SHEET]
    assert ws["I2"].data_type == "n" and ws["I2"].value == 1200
    assert ws["I3"].value is None
    assert store.get(2).data.amount is None


def test_save_makes_backup_and_leaves_no_temp_file(tmp_path):
    store = make_store(tmp_path)
    add(store)
    add(store, data=ticket(name="第二張"))
    names = sorted(p.name for p in store.path.parent.iterdir())
    assert names == ["tickets.xlsx", "tickets.xlsx.bak"]
    restored = tmp_path / "restored.xlsx"  # openpyxl 不肯直接開 .bak，先複製成 .xlsx 再檢查
    restored.write_bytes(store.path.with_name("tickets.xlsx.bak").read_bytes())
    assert load_workbook(restored)[RECORD_SHEET].max_row == 2  # 備份是上一次存檔的狀態（只有第一筆）


def test_failed_replace_keeps_original_and_cleans_temp(tmp_path, monkeypatch):
    store = make_store(tmp_path)
    add(store)
    before = store.path.read_bytes()

    def boom(*_a, **_k):
        raise PermissionError("locked")

    monkeypatch.setattr(os, "replace", boom)
    with pytest.raises(StoreError, match="寫入 Excel 檔失敗"):
        add(store, data=ticket(name="寫不進去"))
    monkeypatch.undo()
    assert store.path.read_bytes() == before
    assert not store.path.with_name("tickets.xlsx.tmp").exists()
    assert [r.data.name for r in store.all_records()] == ["測試演唱會 2099"]


# ---------- 流水號、狀態、刪除與更正 ----------

def test_serial_is_never_reused_after_delete(tmp_path):
    store = make_store(tmp_path)
    add(store)
    store.set_status(1, STATUS_DELETED)
    assert add(store, data=ticket(name="新的")).serial == 2
    assert [r.serial for r in store.all_records()] == [1, 2]  # 刪除只改狀態，那一列還在
    assert [r.serial for r in store.list_records(UID_A)] == [2]
    assert [r.serial for r in store.list_records(UID_A, include_inactive=True)] == [1, 2]


def test_update_fields_changes_only_allowed_fields(tmp_path):
    store = make_store(tmp_path)
    add(store)
    updated = store.update_fields(1, {"amount": 1500.0, "name": "改過的名稱"})
    assert (updated.data.amount, updated.data.name) == (1500.0, "改過的名稱")
    assert updated.uploader == "小明" and updated.photo_paths == ("photos/a.jpg",)
    for forbidden in ("status", "uploader", "uploader_id", "photo_paths", "recorded_at", "serial"):
        with pytest.raises(StoreError, match="不能更正"):
            store.update_fields(1, {forbidden: "x"})


def test_unknown_serial_raises(tmp_path):
    store = make_store(tmp_path)
    for call in (lambda: store.update_fields(99, {"name": "x"}), lambda: store.set_status(99, STATUS_REMOVED)):
        with pytest.raises(StoreError, match="找不到"):
            call()
    assert store.get(99) is None


def test_reply_message_ids(tmp_path):
    store = make_store(tmp_path)
    add(store, reply="m1")
    store.add_reply_message_id(1, "m2")
    assert store.find_by_reply_message_id("m2").serial == 1
    assert store.find_by_reply_message_id("nope") is None
    assert store.get(1).reply_message_ids == ("m1", "m2")


# ---------- 重複補齊 ----------

def test_merge_fills_blanks_and_keeps_old_values_when_new_is_empty(tmp_path):
    store = make_store(tmp_path)
    blurry = ticket(seat_or_type="", amount=None, currency="", location="")
    add(store, data=blurry, photo="photos/blurry.jpg")
    clear = ticket(location="清楚的場館", other_info="")
    duplicate = store.find_duplicate(UID_A, clear)
    assert duplicate is not None
    merged = store.merge_into(duplicate.serial, clear, "photos/clear.jpg")
    assert merged.data.location == "清楚的場館"
    assert merged.data.amount == 1200.0 and merged.data.seat_or_type == "A 區 12 排"
    assert merged.data.other_info == "票號：TEST-0001"  # 新資料的 other_info 是空的 → 保留舊值
    assert merged.photo_paths == ("photos/clear.jpg", "photos/blurry.jpg")
    assert len(store.all_records()) == 1


def test_duplicate_search_is_per_uploader_and_active_only(tmp_path):
    store = make_store(tmp_path)
    add(store, uid=UID_A)
    assert store.find_duplicate(UID_B, ticket()) is None
    assert store.find_duplicate(UID_A, ticket()).serial == 1
    store.set_status(1, STATUS_DELETED)
    assert store.find_duplicate(UID_A, ticket()) is None


# ---------- 統計表 ----------

def test_stats_sheet_is_rebuilt_from_records(tmp_path):
    store = make_store(tmp_path)
    add(store, uploader="小明", uid=UID_A, data=ticket(amount=3800.0), when=datetime(2026, 10, 1))
    add(store, uploader="小明", uid=UID_A, data=ticket(name="日本", amount=15000.0, currency="JPY"), when=datetime(2026, 9, 5))
    add(store, uploader="小華", uid=UID_B, data=ticket(name="別人的", amount=999.0), when=datetime(2026, 10, 2))
    add(store, uploader="小明", uid=UID_A, data=ticket(name="要刪的", amount=100.0), when=datetime(2026, 10, 3))
    store.set_status(4, STATUS_DELETED)
    rows = {r[0]: r for r in load_workbook(store.path)[STATS_SHEET].iter_rows(min_row=2, values_only=True)}
    a = rows[UID_A]
    assert (a[1], a[2], a[3], a[4], a[5]) == ("小明", 1, "TWD 3,800", 2, "JPY 15,000；TWD 3,800")
    b = rows[UID_B]
    assert (b[2], b[3], b[4], b[5]) == (1, "TWD 999", 1, "TWD 999")
    assert len(rows) == 2


def test_summary_counts_unknown_amounts(tmp_path):
    store = make_store(tmp_path)
    add(store)
    add(store, data=ticket(name="沒金額", amount=None, currency=""))
    summary = store.summary(UID_A)
    assert summary.month.count == 2 and summary.month.unknown_amount_count == 1
    assert summary.month.amounts == {"TWD": 1200.0}


def test_list_records_filters_by_month(tmp_path):
    store = make_store(tmp_path)
    add(store, when=datetime(2026, 9, 30, 23, 59))
    add(store, data=ticket(name="十月"), when=datetime(2026, 10, 1, 0, 0))
    assert [r.data.name for r in store.list_records(UID_A, month="2026-10")] == ["十月"]
    assert len(store.list_records(UID_A)) == 2


# ---------- 防護 4：公式注入 ----------

def inject_everywhere(store, payload):
    data = ticket(ticket_type=payload, name=payload, location=payload, seat_or_type=payload, other_info=payload,
                  currency="TWD")
    return add(store, uploader=payload, data=data, photo=f"photos/{payload}")


@pytest.mark.parametrize("payload", PAYLOADS)
def test_formula_payloads_are_stored_as_plain_text(tmp_path, payload):
    store = make_store(tmp_path)
    record = inject_everywhere(store, payload)

    assert not has_formula(store.path), "工作表 XML 裡不能有任何公式標籤"

    loaded = store.get(record.serial)  # 讀回來的值必須和寫進去的一模一樣
    expected = payload.replace("\r", "\n")  # 唯一的例外：XML 規格會把 CR 換成 LF，和安全無關
    assert (loaded.uploader, loaded.data.name, loaded.data.location, loaded.data.seat_or_type,
            loaded.data.other_info, loaded.data.ticket_type) == (expected,) * 6

    wb = load_workbook(store.path)
    ws = wb[RECORD_SHEET]
    for column in "CDEGHK":  # 上傳者、種類、名稱、地點、座位、其他資訊
        cell = ws[f"{column}2"]
        assert cell.data_type == "s"
        assert cell.quotePrefix is looks_like_formula(payload)
    stats = wb[STATS_SHEET]
    assert stats["B2"].value == expected and stats["B2"].data_type == "s"  # 統計表裡的上傳者名稱也一樣安全


@pytest.mark.parametrize("payload", PAYLOADS)
def test_formula_payloads_stay_safe_through_update_and_merge(tmp_path, payload):
    store = make_store(tmp_path)
    add(store)
    store.update_fields(1, {"name": payload, "location": payload})
    store.merge_into(1, ticket(name=payload, seat_or_type=payload), "photos/b.jpg")
    store.refresh_stats()
    assert not has_formula(store.path)
    assert store.get(1).data.seat_or_type == payload.replace("\r", "\n")


def test_all_payloads_together_still_have_no_formula(tmp_path):
    store = make_store(tmp_path)
    for payload in PAYLOADS:
        inject_everywhere(store, payload)
    assert not has_formula(store.path)
    assert len(store.all_records()) == len(PAYLOADS)


def test_ordinary_text_is_not_marked_with_quote_prefix(tmp_path):
    store = make_store(tmp_path)
    add(store, data=ticket(name="一般名稱 - 日場", other_info="共 2 張；幣別為推測"))
    ws = load_workbook(store.path)[RECORD_SHEET]
    assert ws["E2"].quotePrefix is False and ws["K2"].quotePrefix is False


@pytest.mark.parametrize("text,expected", [
    ("=1+1", True), ("+1", True), ("-1", True), ("@a", True), ("\tx", True), ("\rx", True), (" =x", True),
    ("\n+x", True), ("＝x", True), ("＋x", True), ("普通", False), ("a=b", False), ("", False), ("1+1", False),
])
def test_looks_like_formula(text, expected):
    assert looks_like_formula(text) is expected


def test_write_cell_rejects_non_finite_numbers_and_unknown_types(tmp_path):
    ws = load_workbook(make_store(tmp_path).path)[RECORD_SHEET]
    for bad in (float("nan"), float("inf"), float("-inf")):
        with pytest.raises(StoreError):
            write_cell(ws, 5, 1, bad)
    for bad in ([1], {"a": 1}, b"bytes", object()):
        with pytest.raises(StoreError):
            write_cell(ws, 5, 1, bad)


def test_numbers_are_never_treated_as_text(tmp_path):
    ws = load_workbook(make_store(tmp_path).path)[RECORD_SHEET]
    write_cell(ws, 5, 1, -2)
    write_cell(ws, 5, 2, 3.5)
    assert ws["A5"].data_type == "n" and ws["A5"].value == -2 and ws["B5"].value == 3.5


def test_no_direct_cell_writes_outside_write_cell():
    """除了 write_cell 以外，這個檔案不准有別的地方直接設定儲存格的值。"""
    source = Path(excel_store.__file__).read_text(encoding="utf-8")
    assignments = re.findall(r"^\s*[\w\.\(\)=, ]*\.value\s*=[^=]", source, re.M)
    in_write_cell = source.count("ws.cell(row=row, column=column).value = value") + source.count("cell.value =")
    assert len(assignments) == in_write_cell, assignments
    assert not re.search(r"\b\w*(ws|sheet)\w*\.append\(", source)  # 沒有用工作表的 append 繞過 write_cell


# ---------- 拿掉防護 4 的對照 ----------

def test_without_guard_it_still_runs_but_formulas_get_through(tmp_path, load_without_guard):
    unsafe = load_without_guard("excel_store.py", 4)
    folder = tmp_path / "unsafe"
    folder.mkdir()
    store = unsafe.TicketStore(folder / "tickets.xlsx", clock=lambda: NOW)

    # 1. 仍然能跑（一般內容完全正常）
    store.add_record(uploader="小明", uploader_id=UID_A, data=ticket(), photo_path="photos/a.jpg", recorded_at=NOW)
    assert store.get(1).data.name == "測試演唱會 2099"
    assert not has_formula(store.path)

    # 2. 以 = 開頭的內容會被存成公式（打開 Excel 就會被執行）
    store.add_record(uploader="攻擊者", uploader_id=UID_B, data=ticket(name='=HYPERLINK("http://example.com","點我")'),
                     photo_path="photos/b.jpg", recorded_at=NOW)
    assert has_formula(store.path)
