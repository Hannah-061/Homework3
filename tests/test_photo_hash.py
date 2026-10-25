"""照片 SHA-256 防重的測試：完全相同的照片（同一上傳者、有效紀錄）不會重複記錄。

涵蓋：流程（辨識前攔截、不存照片、不增加累計、回覆綁定）、重啟、刪除後重傳、併發與重試、
Excel 新欄位與舊格式遷移。全部使用合成資料。
"""

from __future__ import annotations

import asyncio
import hashlib
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook

from synthetic import (
    NOW, FakeMessenger, ScriptedAnalyzer, Sleeper, analysis, jpeg_bytes, make_attachment, make_config, make_message,
)
from ticket_agent.app import build_app
from ticket_agent.excel_store import RECORD_HEADERS, RECORD_SHEET, STATS_HEADERS, STATS_SHEET, StoreError, TicketStore
from ticket_agent.models import STATUS_ACTIVE, STATUS_DELETED, STATUS_REMOVED, TicketData
from ticket_agent.queue_worker import PhotoPipeline

UID_A = "100"
UID_B = "200"
ALREADY = "這張照片先前已記錄過 #1，沒有新增"
OTHER_PHOTO = ("COMPLETELY DIFFERENT TICKET", "Row 9  JPY 5000")


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@pytest.fixture
def world(tmp_path):
    class World:
        config = make_config(tmp_path)

        def build(self, analyzer):
            app = build_app(self.config, self.messenger, analyze=analyzer, sleep=self.sleeper, retry_delay_s=0.0)
            app.start()
            self.apps.append(app)
            return app

        def store(self):
            return TicketStore(self.config.excel_path)

        def photos(self):
            return sorted(p.name for p in (self.config.data_dir / "photos").iterdir())

    w = World()
    w.apps, w.messenger, w.sleeper = [], FakeMessenger(), Sleeper()
    yield w
    for app in w.apps:
        app.lock.release()


async def upload(world, app, *, msg_id="m1", author=UID_A, att_id="a1", data=None, when=NOW):
    world.messenger.downloads[att_id] = jpeg_bytes() if data is None else data
    message = make_message(msg_id, author_id=author, attachments=[make_attachment(att_id)], when=when)
    await app.handle_message(message)
    await app.drain()
    return message


def last_text(world) -> str:
    return world.messenger.text_of(world.messenger.sent[-1]["id"])


# ================================================================== 流程

async def test_same_photo_on_another_day_is_not_recorded_again(world):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2", when=NOW + timedelta(days=1))
    records = world.store().all_records()
    assert len(records) == 1 and records[0].status == STATUS_ACTIVE
    assert records[0].photo_sha256 == (sha(jpeg_bytes()),)


async def test_second_upload_skips_claude_saves_no_photo_and_keeps_totals(world):
    analyzer = ScriptedAnalyzer(analysis())
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1")
    photos_before = world.photos()
    await upload(world, app, msg_id="m2", att_id="a2", when=NOW + timedelta(days=1))

    assert analyzer.calls == 1  # 第二次不呼叫辨識
    assert world.photos() == photos_before  # 不保存重複照片
    text = last_text(world)
    assert ALREADY in text
    assert "本月（2026-10）：1 筆" in text and "歷史：1 筆" in text  # 累計沒有增加
    assert "TEST-0001" not in text  # 票號仍不顯示
    assert len(world.store().list_records(UID_A)) == 1


async def test_duplicate_reply_message_is_bound_so_it_can_still_be_corrected(world):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2")
    assert world.store().get(1).reply_message_ids == ("bot-1", "bot-2")
    assert world.store().find_by_reply_message_id("bot-2").serial == 1


async def test_different_photo_is_recorded_normally(world):
    analyzer = ScriptedAnalyzer(analysis(), analysis(name="乙演唱會", seat_or_type="B 區"))
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2", data=jpeg_bytes(lines=OTHER_PHOTO))
    assert analyzer.calls == 2 and len(world.store().all_records()) == 2


async def test_different_bytes_with_same_fields_still_use_the_old_field_heuristic(world):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2", data=jpeg_bytes(lines=OTHER_PHOTO))
    (record,) = world.store().all_records()
    assert len(record.photo_paths) == 2  # 補齊舊的一筆，兩張照片都保留
    assert record.photo_sha256 == (sha(jpeg_bytes(lines=OTHER_PHOTO)), sha(jpeg_bytes()))  # 兩張的雜湊都記錄
    assert "已更新先前那筆 #1" in last_text(world)


async def test_a_photo_that_was_merged_in_is_also_recognised_next_time(world):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2", data=jpeg_bytes(lines=OTHER_PHOTO))
    await upload(world, app, msg_id="m3", att_id="a3", data=jpeg_bytes(lines=OTHER_PHOTO))
    assert len(world.store().all_records()) == 1 and ALREADY in last_text(world)


async def test_the_same_photo_from_a_different_user_is_recorded_separately(world):
    analyzer = ScriptedAnalyzer(analysis())
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1", author=UID_A)
    await upload(world, app, msg_id="m2", att_id="a2", author=UID_B)
    assert analyzer.calls == 2
    assert sorted(r.uploader_id for r in world.store().all_records()) == [UID_A, UID_B]


async def test_non_ticket_photos_are_not_covered(world):
    from synthetic import not_ticket
    analyzer = ScriptedAnalyzer(not_ticket("一杯咖啡"))
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1")
    await upload(world, app, msg_id="m2", att_id="a2")
    assert analyzer.calls == 2 and len(world.photos()) == 2


async def test_failed_recognition_photos_are_not_covered(world):
    from ticket_agent.claude_runner import ClaudeFailed
    analyzer = ScriptedAnalyzer(*[ClaudeFailed("x")] * 3, analysis())
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1")
    assert world.store().all_records() == []
    await upload(world, app, msg_id="m2", att_id="a2")  # 同一張照片，這次辨識成功：要能記錄
    assert len(world.store().all_records()) == 1


async def test_dedupe_still_works_after_the_bot_restarts_with_the_same_data_dir(world):
    first = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, first, msg_id="m1", att_id="a1")
    await first.stop()
    analyzer = ScriptedAnalyzer(analysis())
    second = world.build(analyzer)
    await upload(world, second, msg_id="m2", att_id="a2")
    assert analyzer.calls == 0 and len(world.store().all_records()) == 1 and ALREADY in last_text(world)


# ---------------------------------------------------------------- 刪除、更正移除之後

@pytest.mark.parametrize("status", [STATUS_DELETED, STATUS_REMOVED])
async def test_reupload_after_delete_or_correction_removal_creates_a_new_record(world, status):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    world.store().set_status(1, status)
    before = world.store().get(1)
    await upload(world, app, msg_id="m2", att_id="a2")

    records = world.store().all_records()
    assert [(r.serial, r.status) for r in records] == [(1, status), (2, STATUS_ACTIVE)]
    assert world.store().get(1) == before  # 舊紀錄沒有被覆寫或復原
    assert "已記錄 #2" in last_text(world)


# ================================================================== 併發與重試

async def test_two_identical_photos_sent_at_the_same_time_make_one_record(world):
    analyzer = ScriptedAnalyzer(analysis())
    app = world.build(analyzer)
    world.messenger.downloads.update({"a1": jpeg_bytes(), "a2": jpeg_bytes()})
    await asyncio.gather(
        app.handle_message(make_message("m1", attachments=[make_attachment("a1")])),
        app.handle_message(make_message("m2", attachments=[make_attachment("a2")])),
    )
    await app.drain()
    assert len(world.store().all_records()) == 1 and analyzer.calls == 1 and len(world.photos()) == 1


async def test_one_message_with_two_identical_photos_makes_one_record(world):
    app = world.build(ScriptedAnalyzer(analysis()))
    world.messenger.downloads.update({"a1": jpeg_bytes(), "a2": jpeg_bytes()})
    await app.handle_message(make_message("m1", attachments=[make_attachment("a1"), make_attachment("a2")]))
    await app.drain()
    assert len(world.store().list_records(UID_A)) == 1 and len(world.photos()) == 1


async def test_authoritative_check_inside_the_write_lock_catches_what_the_early_check_missed(world, monkeypatch):
    """辨識前的檢查只是提早攔截；就算它漏掉，寫入鎖內的檢查也不會新增第二筆，而且只清掉這次剛存的照片。"""
    analyzer = ScriptedAnalyzer(analysis())
    app = world.build(analyzer)
    await upload(world, app, msg_id="m1", att_id="a1")
    original_photos = world.photos()

    async def early_check_misses(self, job, photo_sha256):
        return False

    monkeypatch.setattr(PhotoPipeline, "_answer_if_already_recorded", early_check_misses)
    await upload(world, app, msg_id="m2", att_id="a2")

    assert analyzer.calls == 2  # 預檢沒擋到，所以辨識了
    assert len(world.store().all_records()) == 1
    assert world.photos() == original_photos  # 這次剛存的那張被清掉，原本的那張還在
    assert ALREADY in last_text(world)
    assert world.store().get(1).reply_message_ids == ("bot-1", "bot-2")


async def test_cleanup_only_removes_the_file_this_job_created(world, monkeypatch):
    app = world.build(ScriptedAnalyzer(analysis()))
    await upload(world, app, msg_id="m1", att_id="a1")
    outside = world.config.data_dir / "keep-me.txt"
    outside.write_text("x")

    async def early_check_misses(self, job, photo_sha256):
        return False

    monkeypatch.setattr(PhotoPipeline, "_answer_if_already_recorded", early_check_misses)
    # 即使儲存函式回報的路徑在 photos 之外，也不可以被刪
    monkeypatch.setattr("ticket_agent.queue_worker.store_photo", lambda *a, **k: outside)
    await upload(world, app, msg_id="m2", att_id="a2")
    assert outside.exists()
    assert len(world.store().all_records()) == 1


async def test_retry_after_the_row_was_really_written_is_not_mistaken_for_a_duplicate(world, monkeypatch):
    app = world.build(ScriptedAnalyzer(analysis()))
    real_add = TicketStore.add_record
    once = {"done": False}

    def add_then_complain(self, **kwargs):
        record = real_add(self, **kwargs)
        if not once["done"]:
            once["done"] = True
            raise StoreError("寫完之後才出錯")
        return record

    monkeypatch.setattr(TicketStore, "add_record", add_then_complain)
    await upload(world, app)
    assert len(world.store().all_records()) == 1 and len(world.photos()) == 1
    text = last_text(world)
    assert "已記錄 #1" in text and "先前已記錄過" not in text


# ================================================================== Excel 儲存層

def make_store(tmp_path) -> TicketStore:
    folder = tmp_path / "store"
    folder.mkdir()
    return TicketStore(folder / "tickets.xlsx", clock=lambda: NOW)


def ticket(**overrides) -> TicketData:
    base = dict(ticket_type="演唱會", name="測試演唱會 2099", event_datetime="2099-12-31 19:30", seat_or_type="A 區")
    return TicketData(**{**base, **overrides})


def add(store, *, uid=UID_A, digest="", photo="photos/a.jpg", data=None):
    return store.add_record(uploader="小明", uploader_id=uid, data=data or ticket(), photo_path=photo,
                            recorded_at=NOW, photo_sha256=digest)


def test_new_header_is_the_last_column():
    assert len(RECORD_HEADERS) == 16 and RECORD_HEADERS[-1] == "照片SHA256" and RECORD_HEADERS[14] == "回覆訊息ID"


def test_hash_round_trips_through_the_file(tmp_path):
    store = make_store(tmp_path)
    digest = sha(b"abc")
    add(store, digest=digest)
    assert TicketStore(store.path).get(1).photo_sha256 == (digest,)
    assert load_workbook(store.path)[RECORD_SHEET].cell(row=2, column=16).value == digest


def test_find_active_by_hash_is_per_uploader_active_only_and_ignores_empty(tmp_path):
    store = make_store(tmp_path)
    digest = sha(b"abc")
    add(store, digest=digest)
    add(store, digest="", photo="photos/b.jpg", data=ticket(name="別的"))  # 沒有雜湊（例如舊紀錄）
    assert store.find_active_by_hash(UID_A, digest).serial == 1
    assert store.find_active_by_hash(UID_B, digest) is None
    assert store.find_active_by_hash(UID_A, sha(b"other")) is None
    assert store.find_active_by_hash(UID_A, "") is None  # 空白雜湊絕不配對任何紀錄
    store.set_status(1, STATUS_DELETED)
    assert store.find_active_by_hash(UID_A, digest) is None


def test_merge_into_adds_the_new_hash_in_front(tmp_path):
    store = make_store(tmp_path)
    add(store, digest="a" * 64)
    store.merge_into(1, ticket(), "photos/b.jpg", photo_sha256="b" * 64)
    assert store.get(1).photo_sha256 == ("b" * 64, "a" * 64)


LEGACY_HEADERS = (
    "流水號", "記錄時間", "上傳者", "票券種類", "名稱", "日期時間", "地點", "座位或票種", "金額", "幣別",
    "其他資訊", "照片存放位置", "狀態", "上傳者ID", "回覆訊息ID",
)
LEGACY_ROW = [1, "2026-10-24 15:30:12", "小明", "演唱會", "舊票", "2099-01-01", "場館", "A1", 100, "TWD",
              "票號：OLD-1；共 2 張", "photos/old1.jpg\nphotos/old2.jpg", "有效", UID_A, "bot-9,bot-10"]


def write_legacy(path, *, headers=LEGACY_HEADERS, rows=(LEGACY_ROW,), extra=None):
    wb = Workbook()
    ws = wb.active
    ws.title = RECORD_SHEET
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    stats = wb.create_sheet(STATS_SHEET)
    stats.append(list(STATS_HEADERS))
    if extra:
        extra(wb)
    wb.save(path)


def test_legacy_15_column_file_is_migrated_without_losing_anything(tmp_path):
    path = tmp_path / "tickets.xlsx"
    second = [2, "2026-10-25 10:00:00", "小華", "電影", "片名", "", "", "", None, "", "", "photos/x.jpg", "已刪除", UID_B, ""]
    write_legacy(path, rows=(LEGACY_ROW, second))
    original = path.read_bytes()

    store = TicketStore(path, clock=lambda: NOW)

    ws = load_workbook(path)[RECORD_SHEET]
    assert tuple(c.value for c in ws[1]) == RECORD_HEADERS
    first, other = store.all_records()
    assert first.data.name == "舊票" and first.data.other_info == "票號：OLD-1；共 2 張"
    assert first.photo_paths == ("photos/old1.jpg", "photos/old2.jpg") and first.reply_message_ids == ("bot-9", "bot-10")
    assert other.status == STATUS_DELETED and other.data.amount is None
    assert first.photo_sha256 == () and other.photo_sha256 == ()  # 舊紀錄的雜湊保持空白，不補算
    assert [ws.cell(row=r, column=16).value for r in (2, 3)] == [None, None]
    assert (path.with_name("tickets.xlsx.bak")).read_bytes() == original  # 遷移前的備份
    assert not path.with_name("tickets.xlsx.tmp").exists()
    assert store.find_active_by_hash(UID_A, sha(b"anything")) is None


def test_migrated_file_accepts_new_records_and_a_second_start_changes_nothing(tmp_path):
    path = tmp_path / "tickets.xlsx"
    write_legacy(path)
    store = TicketStore(path, clock=lambda: NOW)
    add(store, digest=sha(b"new"), photo="photos/new.jpg", data=ticket(name="新票"))
    snapshot = path.read_bytes()
    TicketStore(path, clock=lambda: NOW)  # 已經是新格式：再啟動不應該再寫檔
    assert path.read_bytes() == snapshot
    assert [r.serial for r in store.all_records()] == [1, 2]


@pytest.mark.parametrize("headers", [
    ("別人的欄位", *LEGACY_HEADERS[1:]),                      # 15 欄但標題不對
    LEGACY_HEADERS[:14],                                      # 少一欄
    (*LEGACY_HEADERS, "別的欄位"),                            # 第 16 欄標題不是照片SHA256
    (*LEGACY_HEADERS[:7], "座位", *LEGACY_HEADERS[8:], "照片SHA256"),  # 16 欄但中間欄位不對
])
def test_files_not_matching_either_format_refuse_to_start_and_stay_untouched(tmp_path, headers):
    path = tmp_path / "tickets.xlsx"
    write_legacy(path, headers=headers, rows=())
    before = path.read_bytes()
    with pytest.raises(StoreError, match="標題列"):
        TicketStore(path)
    assert path.read_bytes() == before and not path.with_name("tickets.xlsx.bak").exists()


def test_legacy_file_with_data_already_in_the_new_column_is_refused_rather_than_overwritten(tmp_path):
    path = tmp_path / "tickets.xlsx"
    write_legacy(path, rows=([*LEGACY_ROW, "有人手動填的東西"],))
    before = path.read_bytes()
    with pytest.raises(StoreError):
        TicketStore(path)
    assert path.read_bytes() == before


def test_migration_failure_keeps_the_original_file(tmp_path, monkeypatch):
    path = tmp_path / "tickets.xlsx"
    write_legacy(path)
    before = path.read_bytes()

    def broken_replace(*_args):
        raise OSError("磁碟錯誤")

    monkeypatch.setattr("ticket_agent.excel_store.os.replace", broken_replace)
    with pytest.raises(StoreError):
        TicketStore(path)
    assert path.read_bytes() == before and not path.with_name("tickets.xlsx.tmp").exists()
