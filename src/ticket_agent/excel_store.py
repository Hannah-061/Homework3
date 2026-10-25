"""Excel 紀錄檔（docs/architecture.md 第 8 節、防護 4）。

- 一個檔、兩張工作表：「紀錄」（明細，只往下加）與「統計」（每位上傳者的累計，由「紀錄」重新算出）。
- 新增 = 讀進整個檔、在最後加一列、存檔；刪除與更正只改那一列的內容，不會移除列，所以流水號不會重複。
- 存檔是先寫暫存檔、留一份 .bak 備份、再原子替換原檔，避免寫到一半當機留下壞檔。
- 所有寫進儲存格的值，一律只經過 write_cell（防護 4）。這個檔案裡不准有別的地方直接設定儲存格的值。
- 一次只能有一個寫入者。多個非同步工作要共用同一把鎖，由呼叫端負責（architecture.md 第 7 節）。
"""

from __future__ import annotations

import math
import os
import shutil
import unicodedata
from collections.abc import Callable, Mapping
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from . import dedupe
from .models import STATUS_ACTIVE, TicketData, TicketRecord
from .totals import Summary, format_amounts, month_key, summarize

RECORD_SHEET = "紀錄"
STATS_SHEET = "統計"
RECORD_HEADERS = (
    "流水號", "記錄時間", "上傳者", "票券種類", "名稱", "日期時間", "地點", "座位或票種", "金額", "幣別",
    "其他資訊", "照片存放位置", "狀態", "上傳者ID", "回覆訊息ID", "照片SHA256",
)
# 加入「照片SHA256」之前的舊格式（前 15 欄）。符合它的檔案啟動時只會多補一個標題，見 _ensure_file
LEGACY_RECORD_HEADERS = RECORD_HEADERS[:-1]
STATS_HEADERS = ("上傳者ID", "上傳者", "當月筆數", "當月金額", "歷史筆數", "歷史金額", "最後更新")
TIME_FORMAT = "%Y-%m-%d %H:%M:%S"
# 「其他資訊」不在裡面：票號收在那個欄位的最前面，一般更正不能整個取代它（見 ticket_number.py）
EDITABLE_FIELDS = ("ticket_type", "name", "event_datetime", "location", "seat_or_type", "amount", "currency")


class StoreError(Exception):
    """Excel 檔讀寫失敗（檔案被占用、毀損、標題列對不上……）。訊息不含使用者的內容。"""


def write_cell(ws: Worksheet, row: int, column: int, value: object) -> None:
    """把一個值寫進儲存格。

    這是「沒有任何防護」的基本版本：以 = 開頭的文字會被 openpyxl 當成公式存起來。
    下面的防護 4 區塊會用安全的版本取代它；單獨拿掉防護 4，程式仍然能跑，只是不再安全。
    """
    ws.cell(row=row, column=column).value = value


# ==== 防護 4：Excel 公式注入防護　開始 ====
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
_LEADING_BLANKS = " \t\r\n\x0b\x0c"


def looks_like_formula(text: str) -> bool:
    """Excel 可能把它當成公式的文字：開頭（忽略前面的空白，並先做全形半形統一）是 = + - @ 或 Tab、CR。"""
    if text.startswith(("\t", "\r")):
        return True
    stripped = unicodedata.normalize("NFKC", text).lstrip(_LEADING_BLANKS)
    return stripped.startswith(_FORMULA_TRIGGERS)


def write_cell(ws: Worksheet, row: int, column: int, value: object) -> None:  # noqa: F811
    """把一個值寫進儲存格。文字一律存成「文字型別」；可能被當公式的文字再加上 Excel 的「文字前綴」。"""
    cell = ws.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = value  # openpyxl 會把 = 開頭的文字標成公式，下一行立刻改回文字
        cell.data_type = "s"
        cell.quotePrefix = looks_like_formula(value)  # 讀回來的值不變，也不會多出看得到的單引號
    elif value is None:
        cell.value = None
    elif isinstance(value, bool):
        cell.value = str(value)
        cell.data_type = "s"
    elif isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise StoreError("不能把非數字或無限大的數值寫進 Excel")
        cell.value = value  # 數字不是文字，不可能變成公式
    else:
        raise StoreError(f"不支援的儲存格內容型別：{type(value).__name__}")
# ==== 防護 4 結束 ====


def _row_values(record: TicketRecord) -> list[object]:
    d = record.data
    return [
        record.serial, record.recorded_at.strftime(TIME_FORMAT), record.uploader, d.ticket_type, d.name,
        d.event_datetime, d.location, d.seat_or_type, d.amount, d.currency, d.other_info,
        "\n".join(record.photo_paths), record.status, record.uploader_id, ",".join(record.reply_message_ids),
        "\n".join(record.photo_sha256),
    ]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _parse_row(values: tuple[object, ...]) -> TicketRecord | None:
    if values[0] in (None, ""):
        return None
    try:
        serial = int(values[0])
        recorded_at = datetime.strptime(_text(values[1]), TIME_FORMAT)
        amount = None if values[8] in (None, "") else float(values[8])
    except (TypeError, ValueError) as error:
        raise StoreError(f"「紀錄」工作表有一列的流水號、時間或金額格式不對（流水號 {values[0]!r}）") from error
    data = TicketData(
        ticket_type=_text(values[3]), name=_text(values[4]), event_datetime=_text(values[5]),
        location=_text(values[6]), seat_or_type=_text(values[7]), amount=amount,
        currency=_text(values[9]), other_info=_text(values[10]),
    )
    return TicketRecord(
        serial=serial, recorded_at=recorded_at, uploader=_text(values[2]), uploader_id=_text(values[13]),
        data=data, photo_paths=tuple(p for p in _text(values[11]).split("\n") if p),
        status=_text(values[12]) or STATUS_ACTIVE,
        reply_message_ids=tuple(i for i in _text(values[14]).split(",") if i),
        photo_sha256=tuple(h for h in _text(values[15]).split("\n") if h),
    )


class TicketStore:
    def __init__(self, path: Path, clock: Callable[[], datetime] = datetime.now):
        self.path = Path(path)
        self._clock = clock
        self._ensure_file()

    # ---------- 讀 ----------

    def all_records(self) -> list[TicketRecord]:
        return self._records_of(self._load())

    def get(self, serial: int) -> TicketRecord | None:
        return next((r for r in self.all_records() if r.serial == serial), None)

    def find_by_reply_message_id(self, message_id: str) -> TicketRecord | None:
        return next((r for r in self.all_records() if message_id in r.reply_message_ids), None)

    def find_by_photo_path(self, photo_path: str) -> TicketRecord | None:
        """用照片位置找紀錄。每張照片的檔名都是唯一的，所以拿它當「這件事做過了沒」的識別，重試才不會重複寫入。"""
        return next((r for r in self.all_records() if photo_path in r.photo_paths), None)

    def find_active_by_hash(self, uploader_id: str, photo_sha256: str) -> TicketRecord | None:
        """同一位上傳者、狀態為「有效」、照片雜湊相同的紀錄（已刪除、已更正移除的不算）。空白雜湊（例如舊紀錄）絕不配對。"""
        digest = photo_sha256.strip().lower()
        if not digest:
            return None
        matches = [r for r in self.all_records()
                   if r.uploader_id == uploader_id and r.status == STATUS_ACTIVE and digest in r.photo_sha256]
        return max(matches, key=lambda r: r.serial) if matches else None

    def list_records(self, uploader_id: str, *, month: str | None = None, include_inactive: bool = False) -> list[TicketRecord]:
        """某位上傳者自己的紀錄，依流水號由小到大。month 是 'YYYY-MM'，None 代表全部。"""
        return [
            r for r in self.all_records()
            if r.uploader_id == uploader_id
            and (include_inactive or r.status == STATUS_ACTIVE)
            and (month is None or month_key(r.recorded_at) == month)
        ]

    def summary(self, uploader_id: str) -> Summary:
        return summarize(self.list_records(uploader_id), self._clock())

    def find_duplicate(self, uploader_id: str, data: TicketData) -> TicketRecord | None:
        return dedupe.find_duplicate(self.all_records(), uploader_id, data)

    # ---------- 寫 ----------

    def add_record(
        self, *, uploader: str, uploader_id: str, data: TicketData, photo_path: str,
        recorded_at: datetime | None = None, reply_message_id: str | None = None, photo_sha256: str = "",
    ) -> TicketRecord:
        wb = self._load()
        records = self._records_of(wb)
        record = TicketRecord(
            serial=max((r.serial for r in records), default=0) + 1,
            recorded_at=recorded_at or self._clock(), uploader=uploader, uploader_id=uploader_id, data=data,
            photo_paths=(photo_path,), reply_message_ids=(reply_message_id,) if reply_message_id else (),
            photo_sha256=(photo_sha256.lower(),) if photo_sha256 else (),
        )
        ws = wb[RECORD_SHEET]
        self._write_row(ws, ws.max_row + 1, record)
        self._finish(wb, records + [record])
        return record

    def merge_into(self, serial: int, data: TicketData, photo_path: str, photo_sha256: str = "") -> TicketRecord:
        """重複上傳：用新資料補齊舊的那一筆，兩張照片都保留（新的在前）。新照片的雜湊也記下來，之後再傳它也認得。"""
        return self._change(serial, lambda r: replace(
            r, data=dedupe.merge_ticket_data(r.data, data), photo_paths=(photo_path, *r.photo_paths),
            photo_sha256=(photo_sha256.lower(), *r.photo_sha256) if photo_sha256 else r.photo_sha256,
        ))

    def update_fields(self, serial: int, changes: Mapping[str, object]) -> TicketRecord:
        """更正：只能改 EDITABLE_FIELDS 裡的欄位（記錄時間、上傳者、照片位置、狀態不能改）。"""
        unknown = set(changes) - set(EDITABLE_FIELDS)
        if unknown:
            raise StoreError(f"這些欄位不能更正：{', '.join(sorted(unknown))}")
        return self._change(serial, lambda r: replace(r, data=replace(r.data, **changes)))

    def set_status(self, serial: int, status: str) -> TicketRecord:
        return self._change(serial, lambda r: replace(r, status=status))

    def add_reply_message_id(self, serial: int, message_id: str) -> TicketRecord:
        return self._change(serial, lambda r: replace(r, reply_message_ids=(*r.reply_message_ids, message_id)))

    def refresh_stats(self) -> None:
        wb = self._load()
        self._finish(wb, self._records_of(wb))

    # ---------- 內部 ----------

    def _change(self, serial: int, transform: Callable[[TicketRecord], TicketRecord]) -> TicketRecord:
        wb = self._load()
        records = self._records_of(wb)
        ws = wb[RECORD_SHEET]
        for index, record in enumerate(records):
            if record.serial == serial:
                updated = transform(record)
                self._write_row(ws, self._row_of(ws, serial), updated)
                records[index] = updated
                self._finish(wb, records)
                return updated
        raise StoreError(f"找不到流水號 {serial}")

    def _row_of(self, ws: Worksheet, serial: int) -> int:
        for row_index, row in enumerate(ws.iter_rows(min_row=2, max_col=1, values_only=True), start=2):
            if row[0] not in (None, "") and int(row[0]) == serial:
                return row_index
        raise StoreError(f"找不到流水號 {serial}")

    def _write_row(self, ws: Worksheet, row_index: int, record: TicketRecord) -> None:
        for column, value in enumerate(_row_values(record), start=1):
            write_cell(ws, row_index, column, value)

    def _records_of(self, wb: Workbook) -> list[TicketRecord]:
        rows = wb[RECORD_SHEET].iter_rows(min_row=2, max_col=len(RECORD_HEADERS), values_only=True)
        return [r for r in (_parse_row(row) for row in rows) if r is not None]

    def _finish(self, wb: Workbook, records: list[TicketRecord]) -> None:
        self._rewrite_stats(wb[STATS_SHEET], records)
        self._save(wb)

    def _rewrite_stats(self, ws: Worksheet, records: list[TicketRecord]) -> None:
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        now = self._clock()
        by_user: dict[str, list[TicketRecord]] = {}
        for record in records:
            by_user.setdefault(record.uploader_id, []).append(record)
        row_index = 2
        for uploader_id in sorted(by_user):
            user_records = by_user[uploader_id]
            summary = summarize(user_records, now)
            latest_name = max(user_records, key=lambda r: r.serial).uploader
            values = [
                uploader_id, latest_name, summary.month.count, format_amounts(summary.month.amounts),
                summary.history.count, format_amounts(summary.history.amounts), now.strftime(TIME_FORMAT),
            ]
            for column, value in enumerate(values, start=1):
                write_cell(ws, row_index, column, value)
            row_index += 1

    def _ensure_file(self) -> None:
        if not self.path.exists():
            wb = Workbook()
            records_ws = wb.active
            records_ws.title = RECORD_SHEET
            stats_ws = wb.create_sheet(STATS_SHEET)
            for ws, headers in ((records_ws, RECORD_HEADERS), (stats_ws, STATS_HEADERS)):
                for column, title in enumerate(headers, start=1):
                    write_cell(ws, 1, column, title)
            self._save(wb)
            return
        wb = self._load()
        legacy = self._is_legacy_record_sheet(wb[RECORD_SHEET])
        if not legacy and tuple(c.value for c in wb[RECORD_SHEET][1][: len(RECORD_HEADERS)]) != RECORD_HEADERS:
            raise StoreError("Excel 檔「紀錄」工作表的標題列和預期的不一樣，為避免寫進別人的檔案，不啟動")
        if tuple(c.value for c in wb[STATS_SHEET][1][: len(STATS_HEADERS)]) != STATS_HEADERS:
            raise StoreError("Excel 檔「統計」工作表的標題列和預期的不一樣，為避免寫進別人的檔案，不啟動")
        if legacy:
            # 舊格式（沒有「照片SHA256」欄）：只補一個標題，舊紀錄的雜湊保持空白；走一般的備份與原子替換
            write_cell(wb[RECORD_SHEET], 1, len(RECORD_HEADERS), RECORD_HEADERS[-1])
            self._save(wb)

    @staticmethod
    def _is_legacy_record_sheet(ws: Worksheet) -> bool:
        """前 15 欄標題完全符合舊格式，而且第 16 欄以後（標題與所有資料列）都是空的，才算可以安全遷移。
        第 16 欄已經有東西（例如有人手動填過）就不是舊格式：交給後面的標題檢查拒絕，不覆蓋任何內容。"""
        legacy_len = len(LEGACY_RECORD_HEADERS)
        if tuple(c.value for c in ws[1][:legacy_len]) != LEGACY_RECORD_HEADERS:
            return False
        if any(c.value not in (None, "") for c in ws[1][legacy_len:]):
            return False
        return all(
            value in (None, "")
            for row in ws.iter_rows(min_row=2, min_col=legacy_len + 1, values_only=True) for value in row
        )

    def _load(self) -> Workbook:
        try:
            wb = load_workbook(self.path)
        except (OSError, BadZipFile, InvalidFileException, KeyError) as error:
            raise StoreError("讀不到或打不開 Excel 檔（可能被占用或毀損）") from error
        if RECORD_SHEET not in wb.sheetnames or STATS_SHEET not in wb.sheetnames:
            raise StoreError(f"Excel 檔缺少「{RECORD_SHEET}」或「{STATS_SHEET}」工作表")
        return wb

    def _save(self, wb: Workbook) -> None:
        """先存暫存檔，留一份備份，再原子替換原檔。檔名都是固定的，不含任何外部輸入。"""
        temp = self.path.with_name(self.path.name + ".tmp")
        backup = self.path.with_name(self.path.name + ".bak")
        try:
            wb.save(temp)
            os.chmod(temp, 0o600)
            if self.path.exists():
                shutil.copy2(self.path, backup)
            os.replace(temp, self.path)
        except OSError as error:
            temp.unlink(missing_ok=True)
            raise StoreError("寫入 Excel 檔失敗（檔案可能正被其他程式占用）") from error
